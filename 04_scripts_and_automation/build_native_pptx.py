#!/usr/bin/env python3
"""
Gerador Nativo de Apresentações PowerPoint (PPTX) - NOVA IMS & Academic Suite.
Gera apresentações 16:9 profissionais compatíveis com MS Office, Keynote e Google Slides.

Utilização:
    uv run --with python-pptx python build_native_pptx.py
    ou
    python3 build_native_pptx.py
"""

import os
import sys

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
except ImportError:
    print("❌ python-pptx não está instalado. Instale com 'pip install python-pptx' ou corra com 'uv run --with python-pptx python build_native_pptx.py'")
    sys.exit(1)

def build_deck(output_path="apresentacao_nativa.pptx"):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]  # Layout em branco

    # Cores Oficiais NOVA IMS & UI
    BLACK = RGBColor(0, 0, 0)
    WHITE = RGBColor(255, 255, 255)
    CARD_BG = RGBColor(17, 17, 17)
    CARD_BORDER = RGBColor(38, 38, 38)
    GREEN = RGBColor(14, 224, 113)       # #0EE071
    BLUE = RGBColor(55, 119, 255)        # #3777FF
    MAGENTA = RGBColor(255, 0, 156)      # #FF009C
    YELLOW = RGBColor(246, 218, 0)       # #F6DA00
    TEXT_MUTED = RGBColor(209, 213, 219) # #D1D5DB
    TEXT_GRAY = RGBColor(156, 163, 175)  # #9CA3AF

    script_dir = os.path.dirname(os.path.abspath(__file__))
    logo_path = os.path.join(script_dir, "..", "03_slidev_starter_template", "public", "ims_logo_negativo.png")

    def add_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BLACK
        bg.line.fill.background()
        return slide

    def add_header(slide, sub_text, main_text, color=GREEN):
        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8.5), Inches(1.2))
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p1 = tf.paragraphs[0]
        p1.text = sub_text.upper()
        p1.font.name = 'Inter'
        p1.font.size = Pt(9)
        p1.font.bold = True
        p1.font.color.rgb = color

        p2 = tf.add_paragraph()
        p2.text = main_text
        p2.font.name = 'Inter'
        p2.font.size = Pt(22)
        p2.font.bold = True
        p2.font.color.rgb = WHITE
        p2.space_before = Pt(4)

        if os.path.exists(logo_path):
            slide.shapes.add_picture(logo_path, Inches(10.5), Inches(0.4), height=Inches(0.45))

    def add_footer(slide, left_txt="Enquadramento do Projeto"):
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.8), Inches(11.733), Inches(0.015))
        line.fill.solid()
        line.fill.fore_color.rgb = CARD_BORDER
        line.line.fill.background()

        tb_l = slide.shapes.add_textbox(Inches(0.8), Inches(6.9), Inches(5.0), Inches(0.4))
        p_l = tb_l.text_frame.paragraphs[0]
        p_l.text = left_txt
        p_l.font.name = 'Inter'
        p_l.font.size = Pt(9)
        p_l.font.color.rgb = TEXT_GRAY

        tb_r = slide.shapes.add_textbox(Inches(7.5), Inches(6.9), Inches(5.033), Inches(0.4))
        p_r = tb_r.text_frame.paragraphs[0]
        p_r.text = "NOVA IMS. DATA WITH PURPOSE"
        p_r.font.name = 'Inter'
        p_r.font.size = Pt(9)
        p_r.font.bold = True
        p_r.font.color.rgb = WHITE
        p_r.alignment = PP_ALIGN.RIGHT

    def create_card(slide, left, top, width, height, border_color=None):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = CARD_BORDER
        card.line.width = Pt(1)

        if border_color:
            bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(0.06), Inches(height))
            bar.fill.solid()
            bar.fill.fore_color.rgb = border_color
            bar.line.fill.background()
        return card

    # =========================================================================
    # SLIDE 1: CAPA
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    add_bg(s1)

    if os.path.exists(logo_path):
        s1.shapes.add_picture(logo_path, Inches(0.8), Inches(0.8), height=Inches(0.55))

    t_box = s1.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.5), Inches(3.8))
    tf1 = t_box.text_frame
    tf1.word_wrap = True

    p_sub = tf1.paragraphs[0]
    p_sub.text = "PROJETO ACADÉMICO / MASTER'S THESIS"
    p_sub.font.name = 'Inter'
    p_sub.font.size = Pt(11)
    p_sub.font.bold = True
    p_sub.font.color.rgb = GREEN

    p_title = tf1.add_paragraph()
    p_title.text = "Otimização Preditiva em Grafos"
    p_title.font.name = 'Inter'
    p_title.font.size = Pt(46)
    p_title.font.bold = True
    p_title.font.color.rgb = WHITE
    p_title.space_before = Pt(8)

    p_desc = tf1.add_paragraph()
    p_desc.text = "Redução Sistemática de Fricção Operacional e Atrasos em Redes Intermodais Urbanas"
    p_desc.font.name = 'Inter'
    p_desc.font.size = Pt(16)
    p_desc.font.color.rgb = TEXT_MUTED
    p_desc.space_before = Pt(14)

    p_author = tf1.add_paragraph()
    p_author.text = "Adriano Vilhena • NOVA Information Management School (NOVA IMS)"
    p_author.font.name = 'Inter'
    p_author.font.size = Pt(11)
    p_author.font.color.rgb = TEXT_GRAY
    p_author.space_before = Pt(28)

    add_footer(s1, "Template de Apresentação PPTX")

    # =========================================================================
    # SLIDE 2: DILEMA E CONTEXTO (ASSIMÉTRICO)
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_bg(s2)
    add_header(s2, "01. Enquadramento Analítico", "A Saturação Centralizada Gera Gargalos em Cascata")

    # Card Esquerda (60%)
    create_card(s2, 0.8, 1.8, 7.2, 4.6, GREEN)
    tb_c1 = s2.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(6.6), Inches(4.2))
    tfc1 = tb_c1.text_frame
    tfc1.word_wrap = True

    p = tfc1.paragraphs[0]
    p.text = "DILEMA EMPÍRICO IDENTIFICADO"
    p.font.name = 'Inter'
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = GREEN

    p = tfc1.add_paragraph()
    p.text = "Sistemas estáticos não reagem a flutuações rápidas de fluxo"
    p.font.name = 'Inter'
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.space_before = Pt(6)

    p = tfc1.add_paragraph()
    p.text = "A incapacidade de redistribuir tráfego antes da formação do estrangulamento gera uma perda operacional estimada em 4.2M€ por ano na AML, com 22 minutos perdidos por passageiro em horários de ponta."
    p.font.name = 'Inter'
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_MUTED
    p.space_before = Pt(12)

    # Card Direita (40%)
    create_card(s2, 8.3, 1.8, 4.2, 4.6, BLUE)
    tb_c2 = s2.shapes.add_textbox(Inches(8.5), Inches(2.0), Inches(3.8), Inches(4.2))
    tfc2 = tb_c2.text_frame
    tfc2.word_wrap = True

    p = tfc2.paragraphs[0]
    p.text = "PERGUNTA & HIPÓTESE"
    p.font.name = 'Inter'
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = BLUE

    p = tfc2.add_paragraph()
    p.text = "Podemos prever anomalias em menos de 60 segundos?"
    p.font.name = 'Inter'
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.space_before = Pt(6)

    p = tfc2.add_paragraph()
    p.text = "A hipótese central testa se a modelação em grafos dinâmicos antecipa bloqueios com ganho superior a 25% face aos modelos estatísticos existentes."
    p.font.name = 'Inter'
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_MUTED
    p.space_before = Pt(10)

    add_footer(s2, "Slide 2 / 3 • Definição do Problema")

    # =========================================================================
    # SLIDE 3: HERO STATISTIC & CONCLUSÃO
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_bg(s3)
    add_header(s3, "02. Resultados Finais", "Redução Comprovada de 34.8% no Tempo de Fricção")

    # Hero Card
    create_card(s3, 0.8, 1.8, 6.0, 4.6, GREEN)
    tb_hero = s3.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.4), Inches(4.2))
    tf_h = tb_hero.text_frame
    tf_h.word_wrap = True

    p = tf_h.paragraphs[0]
    p.text = "GANHO EXPERIMENTAL"
    p.font.name = 'Inter'
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = GREEN

    p = tf_h.add_paragraph()
    p.text = "-34.8%"
    p.font.name = 'Inter'
    p.font.size = Pt(64)
    p.font.bold = True
    p.font.color.rgb = GREEN
    p.space_before = Pt(4)

    p = tf_h.add_paragraph()
    p.text = "Tempo de espera medido no nó intermodal de Campolide e Entrecampos durante o período piloto."
    p.font.name = 'Inter'
    p.font.size = Pt(12)
    p.font.color.rgb = WHITE
    p.space_before = Pt(10)

    # Cards à Direita
    create_card(s3, 7.1, 1.8, 5.4, 2.1, BLUE)
    tb_r1 = s3.shapes.add_textbox(Inches(7.3), Inches(2.0), Inches(5.0), Inches(1.7))
    tfr1 = tb_r1.text_frame
    tfr1.word_wrap = True
    p = tfr1.paragraphs[0]
    p.text = "92.4% DE PRECISÃO DE ALERTA"
    p.font.name = 'Inter'
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p = tfr1.add_paragraph()
    p.text = "Antecipação de incidentes com 20 minutos de margem operacional."
    p.font.name = 'Inter'
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_MUTED
    p.space_before = Pt(4)

    create_card(s3, 7.1, 4.3, 5.4, 2.1, MAGENTA)
    tb_r2 = s3.shapes.add_textbox(Inches(7.3), Inches(4.5), Inches(5.0), Inches(1.7))
    tfr2 = tb_r2.text_frame
    tfr2.word_wrap = True
    p = tfr2.paragraphs[0]
    p.text = "11.6 ms DE LATÊNCIA EDGE"
    p.font.name = 'Inter'
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p = tfr2.add_paragraph()
    p.text = "Compatível com hardware embarcado sem necessidade de GPU em cloud."
    p.font.name = 'Inter'
    p.font.size = Pt(11)
    p.font.color.rgb = TEXT_MUTED
    p.space_before = Pt(4)

    add_footer(s3, "Slide 3 / 3 • Conclusões & Q&A")

    prs.save(output_path)
    print(f"✅ Apresentação gerada com sucesso: {output_path}")

if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "apresentacao_nativa.pptx"
    build_deck(out)
